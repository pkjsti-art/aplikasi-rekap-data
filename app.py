import pandas as pd
import numpy as np
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# 1. Membaca File
clean_filename = 'DATA BERSIH.xlsx'
master_filename = 'MASTER INDIKATOR.xlsx'

df_clean = pd.read_excel(clean_filename)
df_master = pd.read_excel(master_filename)

# 2. Fungsi Normalisasi Teks
def clean_str(val):
    if pd.isna(val):
        return ""
    text = str(val)
    text = re.sub(r'\s+', '', text).upper()
    return text

df_clean['Mesin_Clean'] = df_clean['Mesin'].apply(clean_str)
df_master['clean_var'] = df_master['Nama Var.'].apply(clean_str)

spindle_col_name = None
for col in df_master.columns:
    col_up = str(col).upper()
    if 'SPIND' in col_up or 'SPIN' in col_up:
        spindle_col_name = col
        break
if spindle_col_name is None:
    spindle_col_name = 'SPINDEL'

# 3. Fungsi Pencocokan Master Indikator
def get_master_data(benang_val, mesin_val, col_name):
    b_clean = clean_str(benang_val)
    m_clean = clean_str(mesin_val)

    matched = df_master[df_master['clean_var'] == b_clean]

    if matched.empty:
        return 0.0

    for idx, row in matched.iterrows():
        master_machine_str = str(row['MESIN'])
        machines = [clean_str(m) for m in master_machine_str.split('/')]
        if m_clean in machines or any(m in m_clean or m_clean in m for m in machines):
            try:
                val = float(row[col_name])
                if val > 0:
                    return val
            except:
                pass

    try:
        return float(matched.iloc[0][col_name])
    except:
        return 0.0

# 4. Deteksi Shift
shift_counts = df_clean.groupby(['Tanggal', 'Pegawai', 'Mesin_Clean'])['Shift'].nunique().reset_index()
shift_counts['Status_Shift'] = shift_counts['Shift'].apply(lambda x: 'PANJANG' if x > 1 else 'PENDEK')

df_clean = pd.merge(df_clean, shift_counts[['Tanggal', 'Pegawai', 'Mesin_Clean', 'Status_Shift']], on=['Tanggal', 'Pegawai', 'Mesin_Clean'], how='left')

def determine_shift_final(row):
    shift_val = str(row['Shift']).strip().upper()
    status_val = row['Status_Shift']

    if status_val == 'PENDEK':
        return f"{shift_val} / PENDEK"
    else:
        group_rows = df_clean[
            (df_clean['Tanggal'] == row['Tanggal']) &
            (df_clean['Pegawai'] == row['Pegawai']) &
            (df_clean['Mesin_Clean'] == row['Mesin_Clean'])
        ]
        shifts_in_group = set(group_rows['Shift'].astype(str).str.strip().str.upper())

        if 'SIANG' in shifts_in_group and 'MALAM' in shifts_in_group:
            return "MALAM / PANJANG"
        elif 'PAGI' in shifts_in_group and 'SIANG' in shifts_in_group:
            return "PAGI / PANJANG"
        else:
            return f"{shift_val} / PANJANG"

df_clean['Shift_Final'] = df_clean.apply(determine_shift_final, axis=1)

# 5. Perhitungan Qty Real Cones
group_max_qty = df_clean.groupby(['Tanggal', 'Pegawai'])['Qty Cones'].max().reset_index()
group_max_qty = group_max_qty.rename(columns={'Qty Cones': 'max_qty'})

df_clean = pd.merge(df_clean, group_max_qty, on=['Tanggal', 'Pegawai'], how='left')

real_cones_list = []
for idx, row in df_clean.iterrows():
    q_val = row['Qty Cones']
    max_q = row['max_qty']

    if q_val == max_q:
        real_cones_list.append(round(q_val))
    else:
        b_val = row['Benang']
        m_val = row['Mesin']
        waktu_current = get_master_data(b_val, m_val, 'WAKTU (MENIT)')

        group_rows = df_clean[(df_clean['Tanggal'] == row['Tanggal']) & (df_clean['Pegawai'] == row['Pegawai'])]
        row_max = group_rows[group_rows['Qty Cones'] == max_q].iloc[0]

        waktu_max_qty_item = get_master_data(row_max['Benang'], row_max['Mesin'], 'WAKTU (MENIT)')

        if waktu_max_qty_item > 0:
            real_cones = (q_val * waktu_current) / waktu_max_qty_item
        else:
            real_cones = q_val

        real_cones_list.append(round(real_cones))

df_clean['Qty Real Cones'] = real_cones_list

# 6. Pembuatan Workbook Multi-Sheet Berdasarkan Mesin (Urutan Tab & Sorting Tanggal Terlama ke Terkini)
output_filename = 'rekap_per_mesin.xlsx'
wb = Workbook()
wb.remove(wb.active)

machines_raw = df_clean['Mesin'].unique()

def machine_sort_key(m):
    m_str = str(m).strip().upper()
    custom_order = {
        'SSM 1A': 1, 
        'SSM 1B': 2, 
        'SSM 2A': 3, 
        'SSM 2B': 4, 
        'SSM 3A': 5, 
        'SSM 3B': 6
    }
    for key, rank in custom_order.items():
        if key in m_str:
            return (0, rank, m_str)
    return (1, 99, m_str)

machines = sorted(machines_raw, key=machine_sort_key)

thin_border = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000')
)

header_fill = PatternFill(start_color='9ED3DC', end_color='9ED3DC', fill_type='solid')
color_list = ['BDB2FF', 'FFCEE3', 'FDFFD2', 'CAEDFF']

for mach in machines:
    safe_mach_name = str(mach).strip()
    sheet_name = re.sub(r'[:\\/?*\[\]]', '_', safe_mach_name)[:31]
    if not sheet_name:
        sheet_name = "MESIN"

    ws = wb.create_sheet(title=sheet_name)

    df_mach = df_clean[df_clean['Mesin'] == mach].copy()

    # Header Baris 1
    ws.merge_cells('F1:G1')
    ws.cell(row=1, column=1, value="TANGGAL")
    ws.cell(row=1, column=2, value="NAMA")
    ws.cell(row=1, column=3, value="MESIN")
    ws.cell(row=1, column=4, value="JUMLAH SPINDLE")
    ws.cell(row=1, column=5, value="SHIFT")
    ws.cell(row=1, column=6, value="HASIL SOFTWINDING")
    ws.cell(row=1, column=8, value="HASIL TOTAL\n( CONES )")
    ws.cell(row=1, column=9, value="TARGET 90%\n( CONES )")
    ws.cell(row=1, column=10, value="PERSENTASE")
    ws.cell(row=1, column=11, value="TARGET 100%\n( CONES )")
    ws.cell(row=1, column=12, value="PERSENTASE")
    ws.cell(row=1, column=13, value="CONES")
    ws.cell(row=1, column=14, value="KG")

    # Header Baris 2
    ws.cell(row=2, column=6, value="JENIS BENANG")
    ws.cell(row=2, column=7, value="JUMLAH CONES")

    for col_idx in [1, 2, 3, 4, 5, 8, 9, 10, 11, 12, 13, 14]:
        ws.merge_cells(start_row=1, start_column=col_idx, end_row=2, end_column=col_idx)

    for r in [1, 2]:
        for c in range(1, 15):
            ws.cell(row=r, column=c).fill = header_fill

    current_row = 3

    df_mach['Tanggal_Parsed'] = pd.to_datetime(df_mach['Tanggal'], errors='coerce')
    
    # REVISI UTAMA: Mengurutkan kelompok Tanggal dari yang terlama ke terkini (Ascending / default)
    group_keys = df_mach.sort_values(['Tanggal_Parsed'], ascending=True)[['Tanggal', 'Pegawai', 'Shift']].drop_duplicates().values

    unique_dates = sorted(df_mach['Tanggal_Parsed'].dropna().unique())
    date_to_color_idx = {dt: idx % len(color_list) for idx, dt in enumerate(unique_dates)}

    for dt, peg, shf in group_keys:
        df_group = df_mach[
            (df_mach['Tanggal'] == dt) &
            (df_mach['Pegawai'] == peg) &
            (df_mach['Shift'] == shf)
        ]
        if df_group.empty:
            continue

        start_group_row = current_row

        total_real_cones = int(df_group['Qty Real Cones'].sum()) if df_group['Qty Real Cones'].sum().is_integer() else df_group['Qty Real Cones'].sum()
        shift_status = df_group['Status_Shift'].iloc[0]

        max_row_in_group = df_group.loc[df_group['Qty Cones'].idxmax()]
        ref_benang = max_row_in_group['Benang']
        ref_mesin = max_row_in_group['Mesin']

        spindle_val = get_master_data(ref_benang, ref_mesin, spindle_col_name)
        if spindle_val.is_integer():
            spindle_val = int(spindle_val)

        col_t90 = 'TARGET 90% SHIFT PENDEK' if shift_status == 'PENDEK' else 'TARGET 90% SHIFT PANJANG'
        col_t100 = 'TARGET 100% SHIFT PENDEK' if shift_status == 'PENDEK' else 'TARGET 100% SHIFT PANJANG'

        target_90 = get_master_data(ref_benang, ref_mesin, col_t90)
        if target_90.is_integer():
            target_90 = int(target_90)

        target_100 = get_master_data(ref_benang, ref_mesin, col_t100)
        if target_100.is_integer():
            target_100 = int(target_100)

        for idx, row in df_group.iterrows():
            benang = row['Benang']
            mesin = row['Mesin']
            qty_cones_val = row['Qty Cones']
            if qty_cones_val.is_integer():
                qty_cones_val = int(qty_cones_val)

            kg_val = row['Total'] if 'Total' in row else 0.0

            q_real_val = row['Qty Real Cones']
            if q_real_val.is_integer():
                q_real_val = int(q_real_val)

            ws.cell(row=current_row, column=1, value=str(row['Tanggal'])[:10])
            ws.cell(row=current_row, column=2, value=row['Pegawai'])
            ws.cell(row=current_row, column=3, value=mesin)

            cell_spindle = ws.cell(row=current_row, column=4, value=spindle_val)
            cell_spindle.number_format = '#,##0'

            ws.cell(row=current_row, column=5, value=row['Shift_Final'])
            ws.cell(row=current_row, column=6, value=benang)

            cell_qreal = ws.cell(row=current_row, column=7, value=q_real_val)
            cell_qreal.number_format = '#,##0'

            cell_tot = ws.cell(row=current_row, column=8, value=total_real_cones)
            cell_tot.number_format = '#,##0'

            cell_t90 = ws.cell(row=current_row, column=9, value=target_90)
            cell_t90.number_format = '#,##0' if isinstance(target_90, int) else '#,##0.00'

            p90 = (total_real_cones / target_90) if target_90 > 0 else 0.0
            cell_p90 = ws.cell(row=current_row, column=10, value=p90)
            cell_p90.number_format = '0.0%'

            cell_t100 = ws.cell(row=current_row, column=11, value=target_100)
            cell_t100.number_format = '#,##0' if isinstance(target_100, int) else '#,##0.00'

            p100 = (total_real_cones / target_100) if target_100 > 0 else 0.0
            cell_p100 = ws.cell(row=current_row, column=12, value=p100)
            cell_p100.number_format = '0.0%'

            cell_cones = ws.cell(row=current_row, column=13, value=qty_cones_val)
            cell_cones.number_format = '#,##0'

            cell_kg = ws.cell(row=current_row, column=14, value=kg_val)
            cell_kg.number_format = '#,##0.00'

            current_row += 1

        end_group_row = current_row - 1

        if end_group_row > start_group_row:
            for col_idx in [1, 2, 3, 4, 5, 8, 9, 10, 11, 12]:
                ws.merge_cells(start_row=start_group_row, start_column=col_idx, end_row=end_group_row, end_column=end_col_idx if 'end_col_idx' in locals() else col_idx)

        dt_parsed = pd.to_datetime(dt, errors='coerce')
        if dt_parsed in date_to_color_idx:
            c_idx = date_to_color_idx[dt_parsed]
            hex_color = color_list[c_idx]
            group_fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type='solid')
            for r in range(start_group_row, end_group_row + 1):
                for c in range(1, 15):
                    ws.cell(row=r, column=c).fill = group_fill

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=14):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    min_widths = {1: 14, 2: 18, 3: 12, 4: 15, 5: 18, 6: 25, 7: 14}

    for col_idx in range(1, 15):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[col_letter]:
            val_str = str(cell.value or '')
            lines = val_str.split('\n')
            for line in lines:
                if len(line) > max_len:
                    max_len = len(line)

        calculated_width = max(max_len + 3, 12)
        if col_idx in min_widths:
            calculated_width = max(calculated_width, min_widths[col_idx])

        ws.column_dimensions[col_letter].width = calculated_width

wb.save(output_filename)
print(f"\nFile rekap per mesin berhasil diperbarui dengan sorting Tanggal dari Terlama ke Terkini (Ascending).")

from google.colab import files
files.download(output_filename)
